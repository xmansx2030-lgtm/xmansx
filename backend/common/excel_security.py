"""أمان ملفات Excel المشترك (استيراد الطلاب والمعلمين) — مستخرج في المرحلة 5.

الحماية: امتداد .xlsx فقط (لا xls/xlsm/xlsb) + حد حجم + فحص ZIP ضد القنابل
(عدد المدخلات/الحجم غير المضغوط/بنية xlsx) + قراءة read_only+data_only
(لا تنفيذ صيغ) + حد صفوف.
"""

import zipfile
from pathlib import Path

from django.conf import settings
from openpyxl import load_workbook

from common.errors import ApiError

TOO_LARGE_MESSAGE = "حجم الملف يتجاوز الحد المسموح (10MB)."
UNSUPPORTED_MESSAGE = "صيغة الملف غير مدعومة. المسموح: ملف Excel بامتداد ‎.xlsx فقط."
INVALID_FILE_MESSAGE = "تعذر قراءة الملف. تأكد أنه ملف Excel سليم."
MAX_IMPORT_COLUMNS = 200
MAX_HEADER_SCAN_ROWS = 50


def _bounded_max_column(sheet) -> int:
    """يعالج ملفات الجهات الرسمية التي لا تصرّح بأبعاد الورقة في XML."""
    max_column = sheet.max_column
    if isinstance(max_column, int) and max_column > 0:
        return min(max_column, MAX_IMPORT_COLUMNS)
    return MAX_IMPORT_COLUMNS


def _headers_from_values(values) -> list[str]:
    headers = [str(cell).strip() if cell is not None else "" for cell in values]
    while headers and not headers[-1]:
        headers.pop()
    return headers


def validate_upload(
    uploaded_file,
    *,
    too_large_code: str = "IMPORT_FILE_TOO_LARGE",
    unsupported_code: str = "UNSUPPORTED_IMPORT_FILE",
    invalid_code: str = "INVALID_IMPORT_FILE",
) -> None:
    """فحوص ما قبل الحفظ — ترفع ApiError برموز قابلة للتخصيص لكل نطاق."""
    if uploaded_file.size > settings.STUDENT_IMPORT_MAX_FILE_BYTES:
        raise ApiError(too_large_code, TOO_LARGE_MESSAGE, status_code=400)

    extension = Path(uploaded_file.name).suffix.lower()
    if extension != ".xlsx":  # يشمل رفض .xls و.xlsm (ماكرو) و.xlsb
        raise ApiError(unsupported_code, UNSUPPORTED_MESSAGE, status_code=400)

    _validate_zip_safety(uploaded_file, invalid_code)
    uploaded_file.seek(0)


def _validate_zip_safety(file_obj, invalid_code: str) -> None:
    try:
        with zipfile.ZipFile(file_obj) as archive:
            infos = archive.infolist()
            if len(infos) > settings.STUDENT_IMPORT_MAX_ZIP_ENTRIES:
                raise ApiError(invalid_code, INVALID_FILE_MESSAGE, status_code=400)
            total_uncompressed = sum(info.file_size for info in infos)
            if total_uncompressed > settings.STUDENT_IMPORT_MAX_UNCOMPRESSED_BYTES:
                raise ApiError(invalid_code, INVALID_FILE_MESSAGE, status_code=400)
            names = {info.filename for info in infos}
            if "[Content_Types].xml" not in names:
                raise ApiError(invalid_code, INVALID_FILE_MESSAGE, status_code=400)
    except zipfile.BadZipFile as exc:
        raise ApiError(invalid_code, INVALID_FILE_MESSAGE, status_code=400) from exc
    finally:
        file_obj.seek(0)


def read_headers(file_obj) -> list[str]:
    workbook = load_workbook(file_obj, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        first_row = next(
            sheet.iter_rows(
                min_row=1,
                max_row=1,
                max_col=_bounded_max_column(sheet),
                values_only=True,
            ),
            (),
        )
        return _headers_from_values(first_row)
    finally:
        workbook.close()
        file_obj.seek(0)


def read_header_candidates(
    file_obj, *, max_scan_rows: int = MAX_HEADER_SCAN_ROWS
) -> list[tuple[int, list[str]]]:
    """يعيد الصفوف غير الفارغة المرشحة لترويسة التقرير مع أرقامها الأصلية.

    بعض تقارير الجهات الخارجية تسبق الجدول بشعار وعنوان وعدة صفوف فارغة؛ لذلك
    يختار مستورد المجال الصف الصحيح حسب أسماء أعمدته المعروفة بدل افتراض الصف الأول.
    """
    workbook = load_workbook(file_obj, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        candidates: list[tuple[int, list[str]]] = []
        for row_number, values in enumerate(
            sheet.iter_rows(
                min_row=1,
                max_row=max_scan_rows,
                max_col=_bounded_max_column(sheet),
                values_only=True,
            ),
            start=1,
        ):
            headers = _headers_from_values(values)
            if any(headers):
                candidates.append((row_number, headers))
        return candidates
    finally:
        workbook.close()
        file_obj.seek(0)


def read_rows(
    file_obj,
    *,
    start_row: int = 2,
    too_large_code: str = "IMPORT_FILE_TOO_LARGE",
) -> list[tuple[int, tuple]]:
    workbook = load_workbook(file_obj, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows: list[tuple[int, tuple]] = []
        safe_start_row = max(2, int(start_row))
        for index, values in enumerate(
            sheet.iter_rows(
                min_row=safe_start_row,
                max_col=_bounded_max_column(sheet),
                values_only=True,
            ),
            start=safe_start_row,
        ):
            if values is None or all(v is None or str(v).strip() == "" for v in values):
                continue
            rows.append((index, values))
            if len(rows) > settings.STUDENT_IMPORT_MAX_ROWS:
                raise ApiError(
                    too_large_code,
                    f"عدد الصفوف يتجاوز الحد المسموح ({settings.STUDENT_IMPORT_MAX_ROWS}).",
                    status_code=400,
                )
        return rows
    finally:
        workbook.close()
        file_obj.seek(0)
