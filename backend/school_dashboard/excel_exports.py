"""Excel exports for manager and vice-principal reports."""

from __future__ import annotations

from io import BytesIO
from typing import Any

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Side
from openpyxl.styles.borders import Border
from openpyxl.utils import get_column_letter

from schools.models import School

REPORT_DEFINITIONS: dict[str, dict[str, Any]] = {
    "absence": {
        "title": "تقرير الغياب",
        "sheet": "الغياب",
        "filename": "absence-report.xlsx",
        "columns": [
            ("الطالب", "full_name", 24),
            ("الصف", "grade_name", 14),
            ("الفصل", "section_name", 12),
            ("أيام غياب كامل", "full_absence_days", 16),
            ("أيام غياب جزئي", "partial_absence_days", 16),
            ("حصص غياب", "absent_periods", 14),
            ("حصص بعذر", "excused_absent_periods", 14),
            ("حصص دون عذر", "unexcused_absent_periods", 16),
            ("أيام غير مكتملة", "incomplete_days", 16),
        ],
    },
    "lateness": {
        "title": "تقرير التأخر الصباحي",
        "sheet": "التأخر",
        "filename": "lateness-report.xlsx",
        "columns": [
            ("الطالب", "full_name", 24),
            ("الصف", "grade_name", 14),
            ("الفصل", "section_name", 12),
            ("مرات التأخر الصباحي", "morning_occurrences", 20),
            ("الدقائق المحتسبة", "morning_minutes", 18),
        ],
    },
    "referrals": {
        "title": "تقرير مسار الإحالات",
        "sheet": "الإحالات",
        "filename": "referrals-report.xlsx",
        "columns": [
            ("الطالب", "student.full_name", 24),
            ("الصف", "student.grade_name", 14),
            ("الفصل", "student.section_name", 12),
            ("الفئة", "category_label", 16),
            ("السبب", "reason_label", 22),
            ("المحيل", "created_by_name", 20),
            ("الوكيل المسؤول", "assigned_vice_principal_name", 20),
            ("المرشد", "assigned_counselor_name", 20),
            ("الأولوية", "priority_label", 12),
            ("الحالة", "status_label", 14),
            ("التاريخ", "created_at", 18),
        ],
    },
}


HEADER_FILL = PatternFill("solid", fgColor="0F766E")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Arial", bold=True, size=16, color="102420")
LABEL_FONT = Font(name="Arial", bold=True, color="4B5B58")
BODY_FONT = Font(name="Arial", color="102420")
THIN_BORDER = Border(bottom=Side(style="thin", color="DCE4E2"))


def _nested_value(row: dict, key: str) -> Any:
    current: Any = row
    for part in key.split("."):
        if current is None:
            return ""
        if isinstance(current, dict):
            current = current.get(part)
        else:
            return ""
    return current if current is not None else ""


def _date_label(value: str | None) -> str:
    if not value:
        return ""
    return value.split("T", 1)[0]


def _write_metadata(
    sheet, *, school: School, title: str, context: dict, columns_count: int
) -> None:
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns_count)
    sheet["A1"] = title
    sheet["A1"].font = TITLE_FONT
    sheet["A1"].alignment = Alignment(horizontal="right")

    range_context = context["range"]
    scope = context["scope"]
    scope_label = "كل المدرسة"
    if scope.get("section_id"):
        scope_label = "فصل محدد حسب الفلتر"
    elif scope.get("grade_id"):
        scope_label = "صف محدد حسب الفلتر"

    metadata = [
        ("المدرسة", school.name),
        ("الفترة", f"{range_context['from_date']} - {range_context['to_date']}"),
        ("النطاق", scope_label),
    ]
    for row_index, (label, value) in enumerate(metadata, start=2):
        sheet.cell(row=row_index, column=1, value=label)
        sheet.cell(row=row_index, column=2, value=value)
        sheet.cell(row=row_index, column=1).font = LABEL_FONT
        sheet.cell(row=row_index, column=2).font = BODY_FONT
        sheet.cell(row=row_index, column=1).alignment = Alignment(horizontal="right")
        sheet.cell(row=row_index, column=2).alignment = Alignment(horizontal="right")

    sheet.merge_cells(start_row=5, start_column=1, end_row=5, end_column=columns_count)
    sheet["A5"] = f"عدد النتائج المطابقة: {context.get('count', 0)}"
    sheet["A5"].font = LABEL_FONT
    sheet["A5"].alignment = Alignment(horizontal="right")


def build_report_workbook(*, report_type: str, school: School, payload: dict) -> bytes:
    definition = REPORT_DEFINITIONS[report_type]
    columns = definition["columns"]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = definition["sheet"]
    sheet.sheet_view.rightToLeft = True
    sheet.freeze_panes = "A7"

    context = {**payload["context"], "count": payload["count"]}
    _write_metadata(
        sheet,
        school=school,
        title=definition["title"],
        context=context,
        columns_count=len(columns),
    )

    header_row = 7
    for column_index, (label, _, width) in enumerate(columns, start=1):
        cell = sheet.cell(row=header_row, column=column_index, value=label)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
        cell.border = THIN_BORDER
        sheet.column_dimensions[get_column_letter(column_index)].width = width

    for row_index, report_row in enumerate(payload["results"], start=header_row + 1):
        for column_index, (_, key, _) in enumerate(columns, start=1):
            value = _nested_value(report_row, key)
            if key == "created_at":
                value = _date_label(value)
            cell = sheet.cell(row=row_index, column=column_index, value=value)
            cell.font = BODY_FONT
            cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
            cell.border = THIN_BORDER

    last_column = get_column_letter(len(columns))
    sheet.print_area = f"A1:{last_column}{max(sheet.max_row, header_row)}"

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def workbook_response(*, report_type: str, content: bytes) -> HttpResponse:
    filename = REPORT_DEFINITIONS[report_type]["filename"]
    response = HttpResponse(
        content,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
