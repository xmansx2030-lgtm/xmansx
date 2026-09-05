"""بناء ملفات xlsx حقيقية للاختبارات (openpyxl) — ليست mock dictionaries."""

import io
import re
import zipfile

from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook

NOOR_HEADERS = [
    "رقم الهوية", "اسم الطالب", "الصف", "الفصل",
    "رقم الطالب", "اسم ولي الأمر", "جوال ولي الأمر",
]


def build_xlsx_bytes(rows: list[list], headers: list[str] | None = None) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers if headers is not None else NOOR_HEADERS)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_xlsx_upload(rows: list[list], headers: list[str] | None = None,
                      filename: str = "noor.xlsx") -> SimpleUploadedFile:
    return SimpleUploadedFile(
        filename,
        build_xlsx_bytes(rows, headers),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def build_ministry_staff_upload(rows: list[list], filename: str = "teachers-report.xlsx"):
    """تقرير معلمي المدرسة: تمهيد وزارة التعليم ثم ترويسات في الصف 15 وخلايا مدمجة."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "GetSchoolTeachersDataReport"
    sheet["G3"] = "المملكة العربية السعودية\n"
    sheet["G5"] = " وزارة التعليم"
    sheet["E12"] = "بيانات معلمي المدرسة"
    sheet["C15"] = "الجوال"
    sheet["D15"] = "البريد الإلكتروني"
    sheet["F15"] = "الإسم"
    sheet["H15"] = "رقم الهوية"
    sheet.merge_cells("D15:E15")
    sheet.merge_cells("F15:G15")
    sheet.merge_cells("H15:J15")
    for row_number, (mobile, email, name, national_id) in enumerate(rows, start=16):
        sheet.cell(row_number, 3, mobile)
        sheet.cell(row_number, 4, email)
        sheet.cell(row_number, 6, name)
        sheet.cell(row_number, 8, national_id)
        sheet.merge_cells(start_row=row_number, start_column=4, end_row=row_number, end_column=5)
        sheet.merge_cells(start_row=row_number, start_column=6, end_row=row_number, end_column=7)
        sheet.merge_cells(start_row=row_number, start_column=8, end_row=row_number, end_column=10)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return SimpleUploadedFile(
        filename,
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def build_noor_report_upload(rows: list[list], filename: str = "noor-students-report.xlsx"):
    """تقرير نور واقعي: بيانات تمهيدية ثم عناوين متباعدة في الصف 12."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "بيانات الطلاب"
    sheet["H2"] = "وزارة التعليم"
    sheet["H4"] = "نظام نور"
    sheet["F7"] = "كشف بيانات الطلاب"
    sheet["F9"] = "المدرسة: مدرسة الاختبار"
    sheet["B12"] = "الصف الدراسي"
    sheet["D12"] = "الشعبة"
    sheet["F12"] = "اسم\nالطالب الرباعي"
    sheet["H12"] = "رقم السجل المدني"
    sheet["J12"] = "رقم الطالب"
    sheet["L12"] = "اسم ولي الأمر"
    sheet["N12"] = "جوال ولي الأمر"

    for row_number, row in enumerate(rows, start=13):
        national_id, name, grade, section, number, guardian, mobile = row
        sheet.cell(row_number, 2, grade)
        sheet.cell(row_number, 4, section)
        sheet.cell(row_number, 6, name)
        sheet.cell(row_number, 8, national_id)
        sheet.cell(row_number, 10, number)
        sheet.cell(row_number, 12, guardian)
        sheet.cell(row_number, 14, mobile)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return SimpleUploadedFile(
        filename,
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def build_official_noor_upload(
    pages: list[dict], filename: str = "EL_StudentInfoReport.xlsx"
) -> SimpleUploadedFile:
    """يبني نسخة صناعية بلا PII تطابق تقرير نور الرسمي متعدد الصفحات."""
    workbook = Workbook()
    for page_index, page in enumerate(pages):
        sheet = workbook.active if page_index == 0 else workbook.create_sheet()
        sheet.title = f"Sheet{page_index + 1}"
        offset = 1 if page_index == 0 else 0
        header_row = 20 + offset

        sheet.cell(1 + offset, 27, "المملكة العربية السعودية\n وزارة التعليم")
        sheet.cell(3 + offset, 5, "1447-1448")
        sheet.cell(3 + offset, 11, "العام الدراسي")
        sheet.cell(5 + offset, 5, page["grade"])
        sheet.cell(5 + offset, 12, "الصف")
        sheet.cell(9 + offset, 5, page.get("department", "السنة المشتركة"))
        sheet.cell(9 + offset, 12, "القسم")
        if page.get("section_metadata", True):
            sheet.cell(13 + offset, 5, page["section"])
        sheet.cell(13 + offset, 12, "الفصل")
        sheet.cell(17 + offset, 18, "كشف الطلاب")

        headers = {
            3: "رقم جوال الطالب",
            6: "عنوان القريب",
            9: "اسم قريب الطالب",
            13: "هاتف العمل",
            16: "هاتف المنزل",
            17: "إسم ولي الامر",
            19: "الفصل",
            20: "تاريخ رخصة الاقامة",
            21: "رقم رخصة الاقامة",
            24: "الجنسية",
            25: "تاريخ الميلاد",
            26: "مكان الميلاد",
            28: "حالة القيد",
            29: "اسم الطالب",
            30: "م",
        }
        for column, header in headers.items():
            sheet.cell(header_row, column, header)

        for sequence, row in enumerate(page["rows"], start=1):
            national_id, name, _grade, section, _number, guardian, mobile = row
            row_number = header_row + sequence
            sheet.cell(row_number, 3, mobile or "966500000000")
            sheet.cell(row_number, 17, guardian)
            sheet.cell(row_number, 19, section or page["section"])
            sheet.cell(row_number, 21, national_id)
            sheet.cell(row_number, 28, "مستمر في الدراسة")
            sheet.cell(row_number, 29, name)
            sheet.cell(row_number, 30, sequence)

        # تذييل لا يمثل طالبًا ويجب ألا يدخل في التحليل.
        sheet.cell(header_row + len(page["rows"]) + 2, 1, "نهاية الصفحة")

    raw = io.BytesIO()
    workbook.save(raw)

    # ملفات نور الرسمية لا تحتوي وسم dimension؛ نحاكي ذلك لاختبار القارئ الحقيقي.
    source = io.BytesIO(raw.getvalue())
    output = io.BytesIO()
    with zipfile.ZipFile(source) as input_zip, zipfile.ZipFile(
        output, "w", compression=zipfile.ZIP_DEFLATED
    ) as output_zip:
        for info in input_zip.infolist():
            content = input_zip.read(info.filename)
            if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", info.filename):
                content = re.sub(br"<dimension[^>]*/>", b"", content, count=1)
            output_zip.writestr(info, content)

    return SimpleUploadedFile(
        filename,
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def noor_row(nid: str, name: str, grade: str = "الأول الثانوي", section: str = "1",
             number: str = "", guardian: str = "", mobile: str = "") -> list:
    return [nid, name, grade, section, number, guardian, mobile]
