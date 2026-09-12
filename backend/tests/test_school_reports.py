from datetime import date, datetime, timedelta
from io import BytesIO
from zoneinfo import ZoneInfo

import pytest
from django.utils import timezone
from openpyxl import load_workbook

from academics.models import AcademicYear, AcademicYearStatus
from attendance.models import (
    DailyAbsenceStatus,
    DailyAttendanceSummary,
    DailyCompleteness,
)
from devices.models import ArrivalSource, ArrivalStatus, SchoolArrival
from memberships.models import SchoolMembership
from referrals.models import (
    ReferralCategory,
    ReferralPriority,
    ReferralReason,
    StudentReferral,
)
from students.models import Grade, Section, Student

TODAY = date.today()
TZ = ZoneInfo("Asia/Riyadh")


@pytest.fixture
def report_env(make_school, role_client):
    school = make_school("مدرسة التقارير")
    manager_client, _, _ = role_client(["SCHOOL_MANAGER"], school=school)
    vice_client, _, _ = role_client(["VICE_PRINCIPAL"], school=school)
    year = AcademicYear.objects.create(
        school=school,
        name="2026/2027",
        start_date=TODAY,
        end_date=TODAY + timedelta(days=1),
        status=AcademicYearStatus.ACTIVE,
    )
    grade = Grade.objects.create(school=school, name="الأول", code="G1", sequence=1)
    section = Section.objects.create(school=school, grade=grade, name="أ", code="A")
    student = Student.objects.create(
        school=school,
        national_id_encrypted="encrypted",
        national_id_lookup_hash="a" * 64,
        national_id_masked="******1111",
        full_name="طالب التقرير",
    )
    DailyAttendanceSummary.objects.create(
        school=school,
        student=student,
        academic_year=year,
        section=section,
        attendance_date=TODAY,
        expected_periods=7,
        submitted_periods=7,
        absent_periods=7,
        present_periods=0,
        excused_absent_periods=2,
        unexcused_absent_periods=5,
        completeness_status=DailyCompleteness.COMPLETE,
        absence_status=DailyAbsenceStatus.FULL,
        calculated_at=timezone.now(),
    )
    SchoolArrival.objects.create(
        school=school,
        student=student,
        attendance_date=TODAY,
        first_arrival_at=datetime(TODAY.year, TODAY.month, TODAY.day, 7, 15, tzinfo=TZ),
        raw_late_minutes=15,
        counted_late_minutes=10,
        status=ArrivalStatus.LATE,
        source=ArrivalSource.MANUAL,
    )
    return {
        "school": school,
        "manager": manager_client,
        "vice": vice_client,
        "student": student,
        "grade": grade,
        "section": section,
    }


@pytest.mark.django_db
def test_manager_and_vice_principal_receive_filtered_absence_report(report_env):
    url = (
        f"/api/v1/reports/absence/?preset=TODAY&grade={report_env['grade'].id}"
        "&absence_type=FULL&excuse_type=MIXED"
    )
    for client in (report_env["manager"], report_env["vice"]):
        response = client.get(url)
        assert response.status_code == 200
        assert response.json()["summary"]["students"] == 1
        assert response.json()["results"][0]["full_name"] == "طالب التقرير"
        assert response.json()["results"][0]["unexcused_absent_periods"] == 5


@pytest.mark.django_db
def test_absence_report_exports_real_xlsx(report_env):
    response = report_env["manager"].get("/api/v1/reports/absence/export.xlsx?preset=TODAY")
    assert response.status_code == 200
    assert response["Content-Type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    sheet = workbook["الغياب"]
    assert sheet["A1"].value == "تقرير الغياب"
    assert sheet["A7"].value == "الطالب"
    assert sheet["H7"].value == "حصص دون عذر"
    assert sheet["A8"].value == "طالب التقرير"
    assert sheet["H8"].value == 5


@pytest.mark.django_db
def test_lateness_report_uses_morning_arrivals_only(report_env):
    response = report_env["vice"].get(
        "/api/v1/reports/lateness/?preset=TODAY&min_occurrences=1"
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["summary"] == {
        "students": 1,
        "morning_occurrences": 1,
        "morning_minutes": 10,
    }
    assert payload["results"][0]["morning_occurrences"] == 1


@pytest.mark.django_db
def test_lateness_report_exports_real_xlsx(report_env):
    response = report_env["vice"].get("/api/v1/reports/lateness/export.xlsx?preset=TODAY")
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    sheet = workbook["التأخر"]
    assert sheet["A1"].value == "تقرير التأخر الصباحي"
    assert sheet["D7"].value == "مرات التأخر الصباحي"
    assert sheet["A8"].value == "طالب التقرير"
    assert sheet["D8"].value == 1
    assert sheet["E8"].value == 10


@pytest.mark.django_db
def test_referral_report_filters_priority_and_student(report_env):
    membership = SchoolMembership.objects.get(
        school=report_env["school"], roles__role="SCHOOL_MANAGER"
    )
    referral = StudentReferral.objects.create(
        school=report_env["school"],
        student=report_env["student"],
        source_type="SCHOOL_MANAGER",
        category=ReferralCategory.ATTENDANCE,
        reason_code=ReferralReason.REPEATED_ABSENCE,
        description="غياب متكرر",
        created_by_membership=membership,
        priority=ReferralPriority.HIGH,
    )
    response = report_env["manager"].get(
        f"/api/v1/reports/referrals/?preset=TODAY&priority=HIGH&student={report_env['student'].id}"
    )
    assert response.status_code == 200
    assert response.json()["summary"]["high_priority"] == 1
    assert response.json()["results"][0]["id"] == referral.id


@pytest.mark.django_db
def test_referrals_report_exports_real_xlsx(report_env):
    membership = SchoolMembership.objects.get(
        school=report_env["school"], roles__role="SCHOOL_MANAGER"
    )
    StudentReferral.objects.create(
        school=report_env["school"],
        student=report_env["student"],
        source_type="SCHOOL_MANAGER",
        category=ReferralCategory.ATTENDANCE,
        reason_code=ReferralReason.REPEATED_ABSENCE,
        description="غياب متكرر",
        created_by_membership=membership,
        priority=ReferralPriority.HIGH,
    )
    response = report_env["manager"].get("/api/v1/reports/referrals/export.xlsx?preset=TODAY")
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    sheet = workbook["الإحالات"]
    assert sheet["A1"].value == "تقرير الإحالات للمرشد"
    assert sheet["D7"].value == "الفئة"
    assert sheet["A8"].value == "طالب التقرير"
    assert sheet["D8"].value == "المواظبة"
    assert sheet["H8"].value == "عاجلة"


@pytest.mark.django_db
def test_teacher_cannot_access_school_reports(role_client):
    client, _, _ = role_client(["TEACHER"])
    for path in ("absence", "lateness", "referrals"):
        assert client.get(f"/api/v1/reports/{path}/?preset=TODAY").status_code == 403
        assert client.get(f"/api/v1/reports/{path}/export.xlsx?preset=TODAY").status_code == 403
