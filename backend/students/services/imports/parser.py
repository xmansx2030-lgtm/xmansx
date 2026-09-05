"""قراءة ملفات الطلاب، بما فيها تقرير نور الرسمي متعدد الأوراق.

تقرير نور الرسمي يوزع الفصول على أوراق/صفحات مستقلة، ويضع الصف الدراسي في
رأس الصفحة بدل عمود داخل الجدول. يحول هذا المحلل التقرير إلى صفوف منطقية موحدة
من دون تعديل الملف الأصلي أو نسخ أي بيانات حساسة إلى السجلات.
"""

from dataclasses import dataclass

from django.conf import settings
from openpyxl import load_workbook

from common.errors import ApiError
from common.excel_security import (
    INVALID_FILE_MESSAGE,
    MAX_HEADER_SCAN_ROWS,
    MAX_IMPORT_COLUMNS,
    read_header_candidates,
)
from common.excel_security import read_rows as read_tabular_rows
from common.excel_security import validate_upload as validate_excel_upload
from students.services.imports import mapping as mapping_service

NOOR_OFFICIAL_FORMAT = "NOOR_OFFICIAL_MULTI_SHEET"
SYNTHETIC_GRADE_HEADER = "الصف الدراسي"


def validate_upload(uploaded_file) -> None:
    validate_excel_upload(uploaded_file)


@dataclass(frozen=True)
class SheetLayout:
    sheet_index: int
    sheet_name: str
    header_row: int
    headers: list[str]
    grade: str
    section: str
    detected_rows: int


def _clean_label(value) -> str:
    text = str(value or "").strip()
    text = (
        text.replace("أ", "ا")
        .replace("إ", "ا")
        .replace("آ", "ا")
        .replace("ة", "ه")
        .replace("ـ", "")
    )
    return " ".join(text.split())


def _trim_values(values) -> list:
    result = list(values)
    while result and (result[-1] is None or not str(result[-1]).strip()):
        result.pop()
    return result


def _sheet_header_candidates(sheet) -> list[tuple[int, list[str]]]:
    candidates: list[tuple[int, list[str]]] = []
    for row_number, values in enumerate(
        sheet.iter_rows(
            min_row=1,
            max_row=MAX_HEADER_SCAN_ROWS,
            max_col=MAX_IMPORT_COLUMNS,
            values_only=True,
        ),
        start=1,
    ):
        headers = [str(value).strip() if value is not None else "" for value in values]
        while headers and not headers[-1]:
            headers.pop()
        if any(headers):
            candidates.append((row_number, headers))
    return candidates


def _metadata_value(sheet, label: str, *, before_row: int) -> str:
    if before_row <= 1:
        return ""
    expected = _clean_label(label)
    for values in sheet.iter_rows(
        min_row=1,
        max_row=before_row - 1,
        max_col=MAX_IMPORT_COLUMNS,
        values_only=True,
    ):
        row = list(values)
        label_indexes = [
            index for index, value in enumerate(row) if _clean_label(value) == expected
        ]
        for label_index in label_indexes:
            # تقارير نور RTL: القيمة عادة إلى يسار النقطتين، ثم نجرب اليمين.
            search_indexes = list(range(label_index - 1, -1, -1)) + list(
                range(label_index + 1, len(row))
            )
            for index in search_indexes:
                value = str(row[index] or "").strip()
                if value and value != ":" and _clean_label(value) != expected:
                    return value
    return ""


def _is_student_row(values: tuple, suggested: dict[str, int | None]) -> bool:
    for field in ("full_name", "national_id", "student_number"):
        index = suggested.get(field)
        if index is not None and index < len(values):
            value = values[index]
            if value is not None and str(value).strip():
                return True
    return False


def _detect_sheet_layout(sheet, sheet_index: int) -> SheetLayout | None:
    header_row, headers = mapping_service.select_header_row(
        _sheet_header_candidates(sheet)
    )
    suggested = mapping_service.suggest_mapping(headers)
    grade = _metadata_value(sheet, "الصف", before_row=header_row)
    section = _metadata_value(sheet, "الفصل", before_row=header_row)

    # بصمة تقرير نور الرسمي: الاسم + الهوية/الإقامة + الفصل، والصف في رأس الصفحة.
    if not (
        suggested.get("full_name") is not None
        and suggested.get("national_id") is not None
        and suggested.get("section") is not None
        and grade
    ):
        return None

    detected_rows = sum(
        1
        for values in sheet.iter_rows(
            min_row=header_row + 1,
            max_col=MAX_IMPORT_COLUMNS,
            values_only=True,
        )
        if _is_student_row(values, suggested)
    )
    return SheetLayout(
        sheet_index=sheet_index,
        sheet_name=sheet.title,
        header_row=header_row,
        headers=headers,
        grade=grade,
        section=section,
        detected_rows=detected_rows,
    )


def _official_layouts(workbook) -> list[SheetLayout]:
    return [
        layout
        for index, sheet in enumerate(workbook.worksheets)
        if (layout := _detect_sheet_layout(sheet, index)) is not None
    ]


def analyze_import(file_obj) -> dict:
    """يعيد الرؤوس ونوع التقرير وعدد الأوراق/الطلاب دون حفظ بيانات الطلاب."""
    workbook = load_workbook(file_obj, read_only=True, data_only=True)
    try:
        layouts = _official_layouts(workbook)
        if layouts:
            first = layouts[0]
            headers = list(first.headers)
            suggested = mapping_service.suggest_mapping(headers)
            if suggested.get("grade") is None:
                headers.append(SYNTHETIC_GRADE_HEADER)
            return {
                "header_row": first.header_row,
                "headers": headers,
                "import_format": NOOR_OFFICIAL_FORMAT,
                "source_sheet_count": len(layouts),
                "detected_rows": sum(layout.detected_rows for layout in layouts),
            }
    finally:
        workbook.close()
        file_obj.seek(0)

    header_row, headers = mapping_service.select_header_row(
        read_header_candidates(file_obj)
    )
    return {
        "header_row": header_row,
        "headers": headers,
        "import_format": "TABULAR",
        "source_sheet_count": 1,
        "detected_rows": 0,
    }


def _headers_signature(headers: list[str]) -> tuple[str, ...]:
    return tuple(_clean_label(header) for header in headers)


def _read_official_noor_rows(file_obj) -> list[tuple[int, tuple]]:
    workbook = load_workbook(file_obj, read_only=True, data_only=True)
    try:
        layouts = _official_layouts(workbook)
        if not layouts:
            raise ApiError("INVALID_IMPORT_FILE", INVALID_FILE_MESSAGE, status_code=400)

        first = layouts[0]
        expected_signature = _headers_signature(first.headers)
        rows: list[tuple[int, tuple]] = []
        display_row = first.header_row + 1

        for layout in layouts:
            if _headers_signature(layout.headers) != expected_signature:
                raise ApiError(
                    "INVALID_IMPORT_FILE",
                    "اختلف ترتيب أعمدة تقرير نور بين أوراق الملف.",
                    status_code=400,
                )

            sheet = workbook.worksheets[layout.sheet_index]
            suggested = mapping_service.suggest_mapping(layout.headers)
            section_index = suggested.get("section")
            base_width = len(layout.headers)

            for source_values in sheet.iter_rows(
                min_row=layout.header_row + 1,
                max_col=MAX_IMPORT_COLUMNS,
                values_only=True,
            ):
                if not _is_student_row(source_values, suggested):
                    continue

                values = _trim_values(source_values[:base_width])
                if len(values) < base_width:
                    values.extend([None] * (base_width - len(values)))
                if (
                    section_index is not None
                    and section_index < len(values)
                    and (values[section_index] is None or not str(values[section_index]).strip())
                    and layout.section
                ):
                    values[section_index] = layout.section
                values.append(layout.grade)
                rows.append((display_row, tuple(values)))
                display_row += 1

                if len(rows) > settings.STUDENT_IMPORT_MAX_ROWS:
                    raise ApiError(
                        "IMPORT_FILE_TOO_LARGE",
                        "عدد الصفوف يتجاوز الحد المسموح "
                        f"({settings.STUDENT_IMPORT_MAX_ROWS}).",
                        status_code=400,
                    )
        return rows
    finally:
        workbook.close()
        file_obj.seek(0)


def read_rows(
    file_obj,
    *,
    start_row: int = 2,
    import_format: str = "TABULAR",
) -> list[tuple[int, tuple]]:
    if import_format == NOOR_OFFICIAL_FORMAT:
        return _read_official_noor_rows(file_obj)
    return read_tabular_rows(file_obj, start_row=start_row)
